# -*- coding: utf-8 -*-
"""兑现申请单据查询 —— TOF 接口直连版（替代 tof_cash_export.py 的浏览器 RPA）。

流程：表单 POST /api/settlement/v1/{t}/feeCashMain/getPage 翻页拉全量
      → openpyxl 写 xlsx 到本地。
不点页面上的"导出"按钮（该按钮在调查手册里标记为写操作），数据全部来自查询接口，
天然绕开 ZIP 下载/解压链路；单次请求秒级返回，无需等待浏览器渲染。

默认口径与原 RPA 的 FILTERS 对齐（四项筛选全部生效，已实测 2026-10-27）：
  法人     companyCode=0528     TCL智能科技（宁波）有限公司
  审批状态 processStatusCodes=30 审批中
  产品线   productLineCode=PL002  TCL_空调产品线
  营销组织 salesOrgCode=A001040000 华南大区

手动运行可用参数覆盖：
  python cash_query.py --start-date 2026-08-01 --end-date 2026-10-27
  python cash_query.py --process-status-codes 30,40   # 多状态逗号分隔
  python cash_query.py --dry-run                      # 只打印最终请求，不发

已知边界（详见同目录 README.md）：
  * 审批状态筛选用 processStatusCodes（复数，多个值用英文逗号分隔）；
    单数候选 processStatus / approveStatus 等实测均被服务端静默忽略；
  * 翻页拉到 len(rows) == totalElements 为止；拉不全直接非零退出，不写半截 xlsx；
  * xlsx 只落本地，不发送飞书（发送由平台侧负责，见 README「飞书口径」）。
"""
import argparse
import math
import sys
import time
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from client import AuthExpired, GATEWAY, TENANT_ID, TofClient, TofError

ENDPOINT = '/api/settlement/v1/{t}/feeCashMain/getPage'


def build_url(path):
    """与 TofClient.url 相同的占位符替换，供 --dry-run 免凭据使用。"""
    for ph in ('{t}', '{p}'):
        if ph in path:
            path = path.replace(ph, TENANT_ID, 1)
            break
    return GATEWAY + path

# (xlsx 列名, 响应字段)
COLUMNS = [
    ('单据编号', 'docNumber'),
    ('标题', 'docSubject'),
    ('法人', 'companyName'),
    ('审批状态', 'processStatusName'),
    ('结算客户', 'settleCustName'),
    ('申请金额', 'applyMoney'),
    ('审批金额', 'agreeMoney'),
    ('创建时间', 'creationDate'),
    ('创建人', 'creatorName'),
    ('营销组织', 'salesOrgName'),
    ('当前处理人', 'currentProcessor'),
    ('单据ID', 'id'),
]

# 默认口径 = 原 RPA tof_cash_export.py 的 FILTERS（四项），各编码均已实测生效
DEFAULT_COMPANY_CODE = '0528'
DEFAULT_COMPANY_NAME = 'TCL智能科技（宁波）有限公司'
DEFAULT_PRODUCT_LINE_CODE = 'PL002'
DEFAULT_PRODUCT_LINE_NAME = 'TCL_空调产品线'
DEFAULT_SALES_ORG_CODE = 'A001040000'
DEFAULT_SALES_ORG_NAME = '华南大区'
DEFAULT_PROCESS_STATUS_CODES = '30'
# 页面 ASSIGN_STATUS 枚举（minified bundle 反推 + 实测 '30' 生效）
STATUS_NAMES = {'10': '已驳回', '20': '编辑中', '30': '审批中', '40': '已审批', '50': '已取消'}

DEFAULT_WINDOW_DAYS = 89     # 默认取近 90 天（与平台调用口径一致）
MAX_SPAN_DAYS = 366          # 超过一年给出提示，但不阻断（服务端实测不限制跨度）
PAGE_SIZE = 500              # 服务端接受 500/1000，实测返回条数与 totalElements 一致
MAX_PAGES = 200              # 安全阀：200 × 500 = 10 万条，超了说明筛选有问题


class IncompleteData(TofError):
    """翻页未拉全（遇到 --max-pages 安全阀或服务端提前断流）。"""


def parse_args(argv=None):
    today = date.today()
    p = argparse.ArgumentParser(description='兑现申请单据查询（TOF 接口直连）')
    p.add_argument('--company-code', default=DEFAULT_COMPANY_CODE,
                   help=f'法人编码（默认 {DEFAULT_COMPANY_CODE} = {DEFAULT_COMPANY_NAME}）')
    p.add_argument('--product-line-code', default=DEFAULT_PRODUCT_LINE_CODE,
                   help=f'产品线编码（默认 {DEFAULT_PRODUCT_LINE_CODE} = {DEFAULT_PRODUCT_LINE_NAME}）')
    p.add_argument('--sales-org-code', default=DEFAULT_SALES_ORG_CODE,
                   help=f'营销组织编码（默认 {DEFAULT_SALES_ORG_CODE} = {DEFAULT_SALES_ORG_NAME}）')
    p.add_argument('--process-status-codes', default=DEFAULT_PROCESS_STATUS_CODES,
                   help='审批状态编码，多个用逗号分隔：10 已驳回 / 20 编辑中 / 30 审批中 / '
                        f'40 已审批 / 50 已取消（默认 {DEFAULT_PROCESS_STATUS_CODES} = 审批中）')
    p.add_argument('--no-process-status', action='store_true',
                   help='不筛审批状态（导出全部状态；等价于 --process-status-codes 空）')
    p.add_argument('--start-date', '--start', dest='start_date',
                   default=(today - timedelta(days=DEFAULT_WINDOW_DAYS)).isoformat(),
                   help='填单日期起 YYYY-MM-DD（默认近 90 天；--start 为兼容别名）')
    p.add_argument('--end-date', '--end', dest='end_date', default=today.isoformat(),
                   help='填单日期止 YYYY-MM-DD（默认今天；--end 为兼容别名）')
    p.add_argument('--max-pages', type=int, default=MAX_PAGES,
                   help=f'翻页安全阀（默认 {MAX_PAGES} 页 × {PAGE_SIZE} 条；'
                        '触顶说明未拉全，脚本会失败而不是写半截文件）')
    p.add_argument('--out', default='',
                   help='输出 xlsx 路径（默认 兑现申请单据查询_<时间戳>.xlsx）')
    p.add_argument('--auth-file', default='',
                   help='auth.json 路径（默认读环境变量 TOF_AUTH_JSON / TOF_BEARER_TOKEN）')
    p.add_argument('--no-renew', action='store_true',
                   help='token 过期时不自动续期，直接失败（默认开浏览器免登续期后重试一次）')
    p.add_argument('--renew-headful', action='store_true',
                   help='续期时弹出可见浏览器；SSO 登录态也过期时，用它在弹出的窗口手动登录一次')
    p.add_argument('--dry-run', action='store_true',
                   help='只打印最终生效的全部筛选条件，不发请求')
    # 平台约定参数：发送由平台侧统一负责，脚本侧接受但忽略（见 README「飞书口径」）
    p.add_argument('--no-lark', action='store_true', help=argparse.SUPPRESS)
    args = p.parse_args(argv)
    if args.no_process_status:
        args.process_status_codes = ''
    return args


def validate_dates(start: date, end: date):
    if start > end:
        raise SystemExit(f'[参数错误] --start-date {start} 晚于 --end-date {end}')
    span = (end - start).days
    if span > MAX_SPAN_DAYS:
        print(f'[提示] 日期跨度 {span} 天超过一年，结果量可能很大；'
              '服务端实测不限制跨度，但建议缩小区间', file=sys.stderr)


def build_form(args, page=0):
    """把生效的筛选条件拼成表单；空值不传（服务端会把空串当"不筛"）。"""
    form = {
        'startDate': args.start_date,
        'endDate': args.end_date,
        'page': page,
        'size': PAGE_SIZE,
    }
    if args.company_code:
        form['companyCode'] = args.company_code
    if args.product_line_code:
        form['productLineCode'] = args.product_line_code
    if args.sales_org_code:
        form['salesOrgCode'] = args.sales_org_code
    if args.process_status_codes:
        form['processStatusCodes'] = args.process_status_codes
    return form


def describe_filters(args):
    """--dry-run / 运行日志用：把最终生效的全部筛选条件翻成人话。"""
    status = args.process_status_codes or '（不筛）'
    if args.process_status_codes:
        names = [STATUS_NAMES.get(c.strip(), c.strip())
                 for c in args.process_status_codes.split(',') if c.strip()]
        status = f'{args.process_status_codes}（{"、".join(names)}）'
    return [
        f'法人        companyCode={args.company_code or "（不筛）"}'
        + (f'（{DEFAULT_COMPANY_NAME}）' if args.company_code == DEFAULT_COMPANY_CODE else ''),
        f'产品线      productLineCode={args.product_line_code or "（不筛）"}'
        + (f'（{DEFAULT_PRODUCT_LINE_NAME}）'
           if args.product_line_code == DEFAULT_PRODUCT_LINE_CODE else ''),
        f'营销组织    salesOrgCode={args.sales_org_code or "（不筛）"}'
        + (f'（{DEFAULT_SALES_ORG_NAME}）'
           if args.sales_org_code == DEFAULT_SALES_ORG_CODE else ''),
        f'审批状态    processStatusCodes={status}',
        f'填单日期    {args.start_date} ~ {args.end_date}',
    ]


def fetch_all(client, form, max_pages):
    """翻页拉全量：一直到 len(rows) == totalElements。返回 (rows, total)。

    进度写 stderr，stdout 保持纯数据。拉不全抛 IncompleteData（调用方不写文件）。
    """
    rows, total = [], None
    for page in range(max_pages):
        data = client.post_form(ENDPOINT, {**form, 'page': page})
        if not isinstance(data, dict) or 'content' not in data:
            raise TofError(f'响应不是分页结构（第 {page} 页）：{str(data)[:200]}')
        if total is None:
            total = data.get('totalElements') or 0
        batch = data.get('content') or []
        rows.extend(batch)
        print(f'[page {page}] 本页 {len(batch)} 条 / 累计 {len(rows)} / 服务端总数 {total}',
              file=sys.stderr)
        if not batch or len(rows) >= total:
            break
    if total is None or len(rows) < total:
        raise IncompleteData(
            f'未拉全：已拉 {len(rows)} / 总数 {total}，'
            f'触到 --max-pages={max_pages} 安全阀或服务端提前断流；'
            '不写半截 xlsx，请缩小日期区间或调大 --max-pages 后重试')
    return rows, total


def check_status_filter(rows, codes):
    """自检审批状态筛选是否真的生效：全部不匹配就报错，避免静默出错数。"""
    if not codes or not rows:
        return
    expected = {STATUS_NAMES[c.strip()] for c in codes.split(',')
                if c.strip() in STATUS_NAMES}
    if not expected:
        return
    matched = sum(1 for row in rows if row.get('processStatusName') in expected)
    ratio = matched / len(rows)
    print(f'[自检] 审批状态命中 {matched}/{len(rows)}（期望 {sorted(expected)}）',
          file=sys.stderr)
    if matched == 0:
        raise TofError(
            f'审批状态筛选未生效：{len(rows)} 条结果没有一条命中 {sorted(expected)}，'
            f'实际状态示例 {sorted({str(r.get("processStatusName")) for r in rows[:5]})}；'
            '请核对 processStatusCodes 参数是否被服务端忽略')


def write_xlsx(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = '兑现申请单据'
    ws.append([name for name, _ in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(field) for _, field in COLUMNS])
    for idx, (name, _) in enumerate(COLUMNS, start=1):
        width = max(10, min(60, len(name) * 2 + 4))
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    wb.save(out_path)


def _fetch_with_renew(args, form):
    """拉数 + token 过期自动续期：AuthExpired 时免登续期一次再重试；
    重试仍过期（SSO 登录态也掉了）→ AuthExpired 冒泡，按退出码 3 报错，
    提示用 --renew-headful 手动登录一次。"""

    def once():
        client = TofClient(auth_file=args.auth_file or None)
        return fetch_all(client, form, args.max_pages)

    try:
        return once()
    except AuthExpired:
        if args.no_renew:
            raise
        import renew
        if not renew.available():
            raise SystemExit('[认证过期] 本机未安装 playwright，无法自动续期；'
                             '请安装后重试，或手工更新 TOF_BEARER_TOKEN')
        mode = '可见浏览器' if args.renew_headful else '后台浏览器'
        print(f'[token] 已过期，自动续期中（{mode}，免登）…', file=sys.stderr)
        out = renew.renew_auth(auth_file=args.auth_file or None,
                               headless=not args.renew_headful)
        print(f'[token] 续期成功 → {out}，重试查询', file=sys.stderr)
        return once()


def main(argv=None):
    args = parse_args(argv)
    start = date.fromisoformat(args.start_date)
    end = date.fromisoformat(args.end_date)
    validate_dates(start, end)
    form = build_form(args)

    if args.dry_run:
        print(f'POST {build_url(ENDPOINT)}')
        print('最终生效的筛选条件：')
        for line in describe_filters(args):
            print(f'  {line}')
        print(f'  表单字段    {sorted(form)}（page/size 由翻页填充，size={PAGE_SIZE}）')
        return 0

    print('[参数] 最终生效的筛选条件：', file=sys.stderr)
    for line in describe_filters(args):
        print(f'  {line}', file=sys.stderr)

    t0 = time.time()
    rows, total = _fetch_with_renew(args, form)
    check_status_filter(rows, args.process_status_codes)

    out = Path(args.out) if args.out else Path(
        f'兑现申请单据查询_{end.strftime("%Y%m%d")}_{time.strftime("%H%M%S")}.xlsx')
    write_xlsx(rows, out)
    pages = math.ceil(total / PAGE_SIZE) if total else 0
    print(f'完成：{len(rows)} 条（服务端总数 {total}，{pages} 页）→ {out.resolve()}'
          f'（耗时 {time.time() - t0:.1f}s）')
    print('文件已落本地；飞书发送由平台侧负责（见 tof_api/README.md「飞书口径」）。')
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except AuthExpired as e:
        print(f'[认证失效] {e}', file=sys.stderr)
        sys.exit(3)
    except IncompleteData as e:
        print(f'[数据不全] {e}', file=sys.stderr)
        sys.exit(5)
    except TofError as e:
        print(f'[接口错误] {e}', file=sys.stderr)
        sys.exit(4)
